"""Compara, celda a celda, la ESTRUCTURA de dos Excel: la plantilla que se mando a VIVE
(privado/, fuera de git) y el que genera la plataforma. No compara los datos (nombres,
dorsales, tallas concretas), solo el formato: hoja, cabeceras, anchos de columna,
celdas con o sin valor, tipo de valor, fuente, relleno, bordes, alineacion y formato.

    python comparar_excel.py <plantilla.xlsx> <generado.xlsx>
"""
import sys
import openpyxl

COLS = "ABCDEFGHIJKLMN"
NOMBRE, DORSAL, TALLA = set("AHL"), set("BEI"), set("CFJM")


def color(c):
    if c is None:
        return None
    if c.type == "rgb":
        return ("rgb", c.rgb)
    if c.type == "theme":
        return ("theme", c.theme, round(c.tint or 0, 3))
    return (c.type, c.value)


def firma(celda, fila):
    f, fl, b, a = celda.font, celda.fill, celda.border, celda.alignment
    col = celda.column_letter
    v = celda.value
    if fila == 1:
        valor = repr(v)
    elif v is None:
        valor = "vacia"
    elif col in NOMBRE:
        valor = "texto" if isinstance(v, str) else f"TIPO {type(v).__name__}"
    elif col in DORSAL:
        valor = "entero" if isinstance(v, int) else f"TIPO {type(v).__name__}"
    elif col in TALLA:
        valor = "talla" if isinstance(v, (int, str)) else f"TIPO {type(v).__name__}"
    else:
        valor = f"valor {type(v).__name__}"
    return {
        "valor": valor,
        "fuente": (f.name, f.sz, bool(f.b), bool(f.i)),
        "relleno": (fl.fill_type, color(fl.fgColor) if fl.fill_type else None),
        "bordes": (b.left.style, b.right.style, b.top.style, b.bottom.style),
        "alineacion": (a.horizontal, a.vertical, bool(a.wrap_text)),
        "formato": celda.number_format,
    }


def main(plantilla, generado):
    wp, wg = openpyxl.load_workbook(plantilla), openpyxl.load_workbook(generado)
    difs = []
    if wp.sheetnames != wg.sheetnames:
        difs.append(f"hojas: {wp.sheetnames} vs {wg.sheetnames}")
    hp, hg = wp.active, wg.active
    if hp.max_row != hg.max_row or hp.max_column != hg.max_column:
        difs.append(f"dimensiones: {hp.dimensions} vs {hg.dimensions}")
    for col in COLS[:13]:
        ap = hp.column_dimensions[col].width if col in hp.column_dimensions else None
        ag = hg.column_dimensions[col].width if col in hg.column_dimensions else None
        if ap is not None and (ag is None or abs(ap - ag) > 0.01):
            difs.append(f"ancho de la columna {col}: {ap} vs {ag}")
    celdas = 0
    for fila in range(1, max(hp.max_row, hg.max_row) + 2):
        for col in COLS:
            celdas += 1
            fp, fg = firma(hp[f"{col}{fila}"], fila), firma(hg[f"{col}{fila}"], fila)
            for k in fp:
                if fp[k] != fg[k]:
                    difs.append(f"{col}{fila} {k}: plantilla={fp[k]} generado={fg[k]}")
    print(f"Celdas comparadas: {celdas} ({COLS[0]}1:{COLS[-1]}{max(hp.max_row, hg.max_row) + 1}), hoja '{hg.title}', dimensiones {hg.dimensions}")
    for d in difs[:40]:
        print("  DIFERENCIA", d)
    print("RESULTADO: IGUAL" if not difs else f"RESULTADO: {len(difs)} DIFERENCIAS")
    return 0 if not difs else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
